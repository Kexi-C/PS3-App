from __future__ import annotations

import html
import importlib.util
import io
import re
import sys
import tempfile
from pathlib import Path

import altair as alt
import numpy as np
import pandas as pd
import inspect

import streamlit as st


# ---------------------------------------------------------------------------
# Paths
# ---------------------------------------------------------------------------
APP_DIR = Path(__file__).resolve().parent


def _resolve_subsystem_root(name: str, aliases: tuple[str, ...] = ()) -> Path:
    """Resolve a subsystem without changing the user's directory layout.

    The packaged layout is ``subsystems/<name>``.  A few local copies keep the
    same subsystem folders directly beside the app, so the canonical path is
    tried first and the fallback is only used when that path is absent.
    """
    names = {name.lower(), *(alias.lower() for alias in aliases)}
    canonical = APP_DIR / "subsystems" / name
    if canonical.is_dir():
        return canonical

    sibling = APP_DIR / name
    if sibling.is_dir():
        return sibling

    # Case-insensitive fallback for macOS/Windows copies.  Restrict the
    # search to the app directory and its immediate parent; never rewrite or
    # create anything in the subsystem tree.
    for base in (APP_DIR, APP_DIR.parent):
        if not base.is_dir():
            continue
        try:
            candidates = sorted(
                path for path in base.rglob("*")
                if path.is_dir()
                and path.name.lower() in names
                and (path / "code").is_dir()
            )
        except OSError:
            candidates = []
        if candidates:
            return candidates[0]
    return canonical


def _code_dir(root: Path) -> Path:
    return root / "code" if (root / "code").is_dir() else root


RAIL_ROOT = _resolve_subsystem_root("rail_corrugation", ("rail", "rail-corrugation"))
DOOR_ROOT = _resolve_subsystem_root("door")
ACV_ROOT = _resolve_subsystem_root("acv", ("acv_leak", "air_conditioning"))
SHM_ROOT = _resolve_subsystem_root("shm", ("structural_health", "structural_health_monitoring"))

RAIL_CODE_DIR = _code_dir(RAIL_ROOT)
RAIL_MODEL_PATH = RAIL_ROOT / "model" / "rail_model.joblib"
DOOR_CODE_DIR = _code_dir(DOOR_ROOT)
ACV_CODE_DIR = _code_dir(ACV_ROOT)
SHM_CODE_DIR = _code_dir(SHM_ROOT)


# ---------------------------------------------------------------------------
# Subsystem definitions (drive the upload guide and the file checks)
# ---------------------------------------------------------------------------
SUBSYSTEMS = {
    "Door": {
        "connected": True,
        "extensions": {".csv"},
        "purpose": "Finds each door opening/closing cycle and flags cycles with abnormal resistance.",
        "upload": "One or more Door stream files such as Train.csv or Test.csv. No answer file is required.",
        "structure": (
            "Train.csv and Test.csv contain continuous door-controller recordings with a timestamp "
            "(e.g. 2023-7-5-0-0-3-760), motor current, voltage, back-EMF, "
            "door timings, commands, switch states and door position. The saved Door model is "
            "loaded from the subsystem for inference."
        ),
        "result": "A list of door cycles, each marked Normal or Abnormal resistance.",
    },
    "ACV": {
        "connected": True,
        "extensions": {".xlsx"},
        "purpose": "Locates the car most likely to have a refrigerant leak in its air-conditioning unit.",
        "upload": "One Excel workbook (.xlsx) per train case. Several cases can be uploaded together.",
        "structure": (
            "First worksheet, headers in the first row: a time column plus per-car "
            "columns named 'Car NN - parameter' (e.g. 'Car 03 - ACV Running Mode'), "
            "one row every 30 seconds."
        ),
        "result": "Every car ranked from most to least likely to be leaking.",
    },
    "Rail Corrugation": {
        "connected": True,
        "extensions": {".csv"},
        "purpose": "Detects wavy wear (corrugation) on the rail surface from axle-box vibration.",
        "upload": "One or more 1-second recordings (.csv), one file per recording.",
        "structure": (
            "129 numeric columns: a speed-sensor pulse (0/1) followed by vibration "
            "and shock for 8 cars × 8 axle boxes, sampled at 10,000 Hz (about 10,000 rows)."
        ),
        "result": "Each recording classed as Normal, Side I corrugation or Side II corrugation.",
    },
    "SHM": {
        "connected": True,
        "extensions": {".csv", ".txt", ".dat", ".npy", ".npz"},
        "purpose": "Estimates accumulated fatigue damage of load-bearing structures from dynamic stress.",
        "upload": "One or more training or test dynamic-stress files. No label file is required.",
        "structure": "The original SHM subsystem reads the uploaded stress files; train/test filenames and file-specific headers are retained.",
        "result": "One cumulative fatigue-damage value per file.",
    },
}

# Full-width keyword compatible with both older and newer Streamlit releases.
FULL = (
    {"width": "stretch"}
    if "width" in inspect.signature(st.button).parameters
    else {"use_container_width": True}
)
# ``st.button`` and ``st.altair_chart`` did not gain the same sizing keyword
# at the same time.  Keep chart sizing independent so the chart call itself
# cannot be rejected by a Streamlit version mismatch.
CHART_FULL = (
    {"width": "stretch"}
    if "width" in inspect.signature(st.altair_chart).parameters
    else {"use_container_width": True}
)

RAIL_FS = 10_000          # sampling frequency, Hz
RAIL_TEETH = 90           # speed-sensor teeth per revolution
RAIL_WHEEL_D = 0.85       # wheel diameter, m
RAIL_N_COLS = 129
RAIL_LABELS = {
    "Normal": "Both rails normal",
    "Side I": "Corrugation suspected on Side I rail",
    "Side II": "Corrugation suspected on Side II rail",
}
RAIL_COLORS = {"Normal": "#2e8b57", "Side I": "#e08a00", "Side II": "#c0392b"}

CAR_COL = re.compile(r"^\s*Car\s*(\d{1,3})\s*-\s*(.+?)\s*$", re.IGNORECASE)
DOOR_TIME = re.compile(r"^\d{4}-\d{1,2}-\d{1,2}-\d{1,2}-\d{1,2}-\d{1,2}-\d{1,3}$")


def _header_token(value) -> str:
    """Normalize CSV headers while preserving the original column names."""
    text = str(value).lstrip("\ufeff").strip().lower()
    return re.sub(r"[^a-z0-9]+", "_", text).strip("_")


# ---------------------------------------------------------------------------
# Page setup
# ---------------------------------------------------------------------------
st.set_page_config(
    page_title="Train Condition Monitoring System",
    page_icon="🚆",
    layout="wide",
    initial_sidebar_state="collapsed",
)

st.markdown(
    """
    <style>
    .block-container { max-width: 1180px; padding-top: 2.2rem; padding-bottom: 3rem; }
    .hero {
        padding: 1.5rem 2rem; border-radius: 18px;
        background: linear-gradient(135deg, #123b63 0%, #1d6fa5 100%);
        color: white; margin-bottom: 1.2rem;
    }
    .hero h1 { margin: 0; font-size: 2rem; font-weight: 700; color: white; }
    .hero p { margin: 0.45rem 0 0; color: #e6f3ff; font-size: 1rem; }
    .section-title { font-size: 1.15rem; font-weight: 650; margin: 0.9rem 0 0.5rem; }
    .guide { padding: 1rem 1.2rem; border: 1px solid #dce7f0; border-radius: 14px;
             background: #f8fbfd; color: #2b3f52; font-size: 0.93rem; line-height: 1.5; }
    .guide b { color: #17324d; }
    .guide p { margin: 0 0 0.55rem; }
    .guide p:last-child { margin-bottom: 0; }
    .verdict { padding: 0.9rem 1.1rem; border-radius: 12px; margin: 0.3rem 0 0.8rem;
               font-size: 1.05rem; font-weight: 600; }
    .verdict.ok { background: #e8f5ee; color: #1e6b43; border-left: 6px solid #2e8b57; }
    .verdict.alert { background: #fdecea; color: #8e2a1f; border-left: 6px solid #c0392b; }
    .verdict.caution { background: #fff5e0; color: #7a4d00; border-left: 6px solid #e08a00; }
    </style>
    """,
    unsafe_allow_html=True,
)

st.session_state.setdefault("uploader_version", {name: 0 for name in SUBSYSTEMS})
st.session_state.setdefault("results", {})


# ---------------------------------------------------------------------------
# File checks (cached per file content)
# ---------------------------------------------------------------------------
def _result(status: str, note: str = "", **features) -> dict:
    return {"status": status, "note": note, "features": features}


@st.cache_data(show_spinner=False, max_entries=1024)
def inspect_rail(key: tuple, _file) -> dict:
    """Lightweight format check: reads only the first rows of the file."""
    data = _file.getvalue()
    try:
        df = pd.read_csv(io.BytesIO(data), header=None, nrows=200, low_memory=False)
    except Exception:  # noqa: BLE001
        return _result("error", "Cannot be read as a CSV file.")
    if df.empty:
        return _result("error", "The file contains no data.")

    has_header = pd.to_numeric(df.iloc[0], errors="coerce").isna().mean() > 0.5
    if has_header:
        df = df.iloc[1:]
    if df.shape[1] != RAIL_N_COLS:
        return _result(
            "error",
            f"Expected 129 columns (speed + 128 vibration/shock channels), found {df.shape[1]}.",
        )
    sample = df.apply(pd.to_numeric, errors="coerce")
    bad_share = float(sample.isna().to_numpy().mean())
    if bad_share > 0.01:
        return _result("error", f"{bad_share:.1%} of the values are empty or non-numeric.")

    notes = []
    n_rows = data.count(b"\n") + (0 if data.endswith(b"\n") else 1) - int(has_header)
    if abs(n_rows - RAIL_FS) > 100:
        notes.append(f"Expected about 10,000 rows (1 s at 10 kHz), found {n_rows:,}.")
    if not sample.iloc[:, 0].dropna().isin([0, 1]).all():
        notes.append("Column 1 is not a 0/1 speed pulse; train speed cannot be estimated.")
    return _result("warn" if notes else "ok", " ".join(notes))


def _read_numeric_csv(data: bytes) -> np.ndarray:
    """Full numeric parse; the pyarrow engine is used when available (faster)."""
    try:
        frame = pd.read_csv(io.BytesIO(data), header=None, engine="pyarrow")
    except Exception:  # noqa: BLE001
        frame = pd.read_csv(io.BytesIO(data), header=None, low_memory=False)
    if pd.to_numeric(frame.iloc[0], errors="coerce").isna().mean() > 0.5:
        frame = frame.iloc[1:]
    return frame.apply(pd.to_numeric, errors="coerce").to_numpy(dtype=float)


@st.cache_data(show_spinner=False, max_entries=1024)
def rail_features(key: tuple, _file) -> dict:
    """Speed estimate and per-car/per-side vibration level, for display only."""
    try:
        values = _read_numeric_csv(_file.getvalue())
    except Exception:  # noqa: BLE001
        return {}
    n_rows = values.shape[0]
    speed_kmh = None
    pulse = values[:, 0][~np.isnan(values[:, 0])]
    if pulse.size and np.isin(np.unique(pulse), [0, 1]).all():
        toggles = int(np.count_nonzero(np.diff(pulse) != 0))
        rev_per_s = toggles / (2 * RAIL_TEETH) / (n_rows / RAIL_FS)
        speed_kmh = rev_per_s * np.pi * RAIL_WHEEL_D * 3.6

    vib = values[:, 1:RAIL_N_COLS:2]  # 64 vibration channels (car-major, position-minor)
    centred = vib - np.nanmean(vib, axis=0)
    rms = np.sqrt(np.nanmean(centred**2, axis=0)).reshape(8, 8)
    side_rms = np.stack([rms[:, 0::2].mean(axis=1), rms[:, 1::2].mean(axis=1)], axis=1)
    return {"speed_kmh": speed_kmh, "side_rms": side_rms.tolist()}


@st.cache_data(show_spinner=False, max_entries=1024)
def inspect_acv(key: tuple, _file) -> dict:
    data = _file.getvalue()
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(min_row=1, max_row=2, values_only=True)
        header = next(rows, ())
        first_data_row = next(rows, None)
        workbook.close()
    except Exception:  # noqa: BLE001
        return _result("error", "Cannot be opened as an Excel workbook.")

    headers = [str(h).strip() for h in header if h is not None]
    cars = sorted({m.group(1) for h in headers if (m := CAR_COL.match(h))})
    if len(cars) < 2:
        return _result(
            "error",
            "The first row of the first worksheet must contain per-car headers such as "
            f"'Car 03 - ACV Running Mode'; found {len(cars)} car(s).",
        )
    if first_data_row is None:
        return _result("error", "The worksheet has headers but no data rows.")

    has_time = any("time" in h.lower() and not CAR_COL.match(h) for h in headers)
    note = "" if has_time else "No time column found in the header row."
    return _result("warn" if note else "ok", note, cars=cars)


@st.cache_data(show_spinner=False, max_entries=1024)
def inspect_door(key: tuple, _file) -> dict:
    """Lightly check a Door Train/Test stream without requiring labels."""
    try:
        data = _file.getvalue()
        df = pd.read_csv(io.BytesIO(data), nrows=200, dtype=str)
    except Exception:  # noqa: BLE001
        return _result("error", "Cannot be read as a CSV file.")
    if df.empty:
        return _result("error", "The file contains no data rows.")

    # The published Door documentation lists 20 stream columns while one
    # section calls them 17.  The original model only needs the named sensor
    # columns, so do not reject a valid file because of that documentation
    # discrepancy.
    if df.shape[1] < 17:
        return _result("error", f"Expected at least 17 Door columns, found {df.shape[1]}.")
    timestamp_column = next(
        (column for column in df.columns if _header_token(column) in {"datetime", "timestamp"}),
        df.columns[0],
    )
    stamps = df[timestamp_column].dropna().str.strip()
    timestamp_share = 0.0 if stamps.empty else float(
        (stamps.str.match(DOOR_TIME) | pd.to_datetime(stamps, errors="coerce").notna()).mean()
    )
    if timestamp_share < 0.5:
        # Some exports are read as headerless by spreadsheet programs.  Check
        # the first column once more without a header before rejecting it.
        try:
            raw = pd.read_csv(io.BytesIO(data), header=None, nrows=200, dtype=str)
            raw_stamps = raw.iloc[:, 0].dropna().str.strip()
            timestamp_share = 0.0 if raw_stamps.empty else float(
                (raw_stamps.str.match(DOOR_TIME)
                 | pd.to_datetime(raw_stamps, errors="coerce").notna()).mean()
            )
        except Exception:  # noqa: BLE001
            pass
    if timestamp_share < 0.5:
        return _result(
            "error",
            "The first column must be a timestamp such as 2023-7-5-0-0-3-760.",
        )
    stem = Path(str(key[1] if len(key) > 1 else "")).stem.lower()
    role = "test" if "test" in stem else "train" if "train" in stem else "stream"
    note = f"Detected as Door {role} input."
    if df.shape[1] not in {17, 20}:
        note += f" {df.shape[1] - 17} extra column(s) beyond the expected 17 are retained."
        return _result("warn", note, role=role)
    return _result("ok", note, role=role)


@st.cache_data(show_spinner=False, max_entries=1024)
def inspect_shm(key: tuple, _file) -> dict:
    # Do not parse SHM files in the UI.  The original SHM reader owns the
    # format (including delimiters, metadata rows and train/test conventions).
    # A pandas pre-check was rejecting valid 6–7 MB files before the subsystem
    # ever received them.  Only reject empty uploads here.
    if not _file.getvalue():
        return _result("error", "The file is empty.")
    return _result("ok", "Accepted; the original SHM subsystem will read this file.")


INSPECTORS = {
    "Door": inspect_door,
    "ACV": inspect_acv,
    "Rail Corrugation": inspect_rail,
    "SHM": inspect_shm,
}


def check_files(subsystem: str, uploaded_files) -> list[dict]:
    spec = SUBSYSTEMS[subsystem]
    allowed = ", ".join(sorted(spec["extensions"]))
    seen: set[str] = set()
    checked = []
    for uploaded in uploaded_files:
        name = Path(uploaded.name).name
        size = uploaded.size
        if name in seen:
            res = _result("error", "Same file name as another uploaded file.")
        elif Path(name).suffix.lower() not in spec["extensions"]:
            res = _result("error", f"{subsystem} accepts {allowed} files only.")
        elif size == 0:
            res = _result("error", "The file is empty.")
        else:
            key = (subsystem, name, size, getattr(uploaded, "file_id", None))
            res = INSPECTORS[subsystem](key, uploaded)
        seen.add(name)
        checked.append(
            {"name": name, "size": size, "file": uploaded,
             "key": (name, size, getattr(uploaded, "file_id", None)), **res}
        )
    return checked


def human_size(n_bytes: int) -> str:
    for unit in ("B", "KB", "MB", "GB"):
        if n_bytes < 1024 or unit == "GB":
            return f"{n_bytes:.0f} {unit}" if unit == "B" else f"{n_bytes:.1f} {unit}"
        n_bytes /= 1024
    return ""


# ---------------------------------------------------------------------------
# Prediction runners (original prediction logic is not modified)
# ---------------------------------------------------------------------------
def _write_inputs(files: list[dict], temp_dir: Path) -> list[Path]:
    paths = []
    for item in files:
        path = temp_dir / item["name"]
        path.write_bytes(item["file"].getvalue())
        paths.append(path)
    return paths


def _module_search_dirs(code_dir: Path, extra_dirs: tuple[Path, ...] = ()) -> list[Path]:
    dirs: list[Path] = []
    for directory in (code_dir, *extra_dirs):
        if directory.is_dir() and directory not in dirs:
            dirs.append(directory)
    if code_dir.name.lower() == "code" and code_dir.parent.is_dir() and code_dir.parent not in dirs:
        dirs.append(code_dir.parent)
    return dirs


def _find_code_file(
    code_dir: Path,
    candidates: tuple[str, ...],
    keywords: tuple[str, ...] = (),
    extra_dirs: tuple[Path, ...] = (),
) -> Path:
    """Find an original subsystem module without copying or changing it."""
    search_dirs = _module_search_dirs(code_dir, extra_dirs)
    # Candidate priority is global: an official predict/adapter file wins over
    # an unrelated helper with a matching keyword in another directory.
    for filename in candidates:
        for directory in search_dirs:
            path = directory / filename
            if path.is_file():
                return path
    if keywords:
        for directory in search_dirs:
            for path in sorted(directory.glob("*.py")):
                stem = path.stem.lower()
                if any(keyword in stem for keyword in keywords):
                    return path
    tried = ", ".join(candidates)
    searched = ", ".join(str(directory) for directory in search_dirs) or str(code_dir)
    raise FileNotFoundError(
        f"No subsystem entrypoint found in {searched} (tried: {tried})."
    )


def _load_original_module(
    code_dir: Path,
    candidates: tuple[str, ...],
    module_key: str,
    keywords: tuple[str, ...] = (),
    extra_dirs: tuple[Path, ...] = (),
):
    """Load one original module by file path, avoiding same-name import collisions."""
    module_path = _find_code_file(code_dir, candidates, keywords, extra_dirs)
    module_name = f"_nebula_{module_key}"
    cached = sys.modules.get(module_name)
    if cached is not None:
        return cached

    code_path = str(code_dir)
    if code_path not in sys.path:
        sys.path.insert(0, code_path)
    spec = importlib.util.spec_from_file_location(module_name, module_path)
    if spec is None or spec.loader is None:
        raise ImportError(f"Cannot load subsystem module: {module_path}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[module_name] = module
    try:
        spec.loader.exec_module(module)
    except Exception:
        sys.modules.pop(module_name, None)
        raise
    return module


def _module_entrypoint(module, names: tuple[str, ...]):
    """Return the first callable adapter entrypoint exposed by a module."""
    for name in names:
        function = getattr(module, name, None)
        if callable(function):
            return function
    for name, function in sorted(vars(module).items()):
        lowered = name.lower()
        if (
            callable(function)
            and not name.startswith("_")
            and any(token in lowered for token in ("predict", "infer", "analy", "damage", "fatigue"))
        ):
            return function
    return None


def _door_low_level_api(module) -> bool:
    return all(
        callable(getattr(module, name, None))
        for name in ("load_stream", "split_segments", "predict_stream")
    ) and callable(getattr(module, "DoorModel", None))


def _invoke_door_stream_entrypoint(function, paths: list[Path]):
    """Invoke a production Door predictor on Train/Test stream files only."""
    try:
        parameters = list(inspect.signature(function).parameters.values())
    except (TypeError, ValueError):
        parameters = []
    positional = [
        parameter for parameter in parameters
        if parameter.kind in (parameter.POSITIONAL_ONLY, parameter.POSITIONAL_OR_KEYWORD)
    ]
    required = [parameter for parameter in positional if parameter.default is parameter.empty]
    if not required:
        return function()

    first = required[0].name.lower()
    batch = (
        first in {"files", "paths", "inputs", "datasets", "streams"}
        or first.endswith("_files")
        or first.endswith("_paths")
        or getattr(function, "__name__", "").lower().endswith(("files", "paths"))
    )
    if batch:
        return function(paths)
    if len(required) == len(paths) and len(paths) > 1:
        return function(*paths)
    if len(paths) == 1:
        return function(paths[0])
    return [function(path) for path in paths]


def _door_result_frame(value) -> pd.DataFrame:
    """Extract a prediction table from common adapter return values."""
    if isinstance(value, tuple):
        for part in value:
            if isinstance(part, (pd.DataFrame, dict, list, tuple, str, Path)):
                try:
                    return _door_result_frame(part)
                except ValueError:
                    continue
    if isinstance(value, pd.DataFrame):
        frame = value.copy()
    elif isinstance(value, dict):
        known = next(
            (value[key] for key in ("submission", "predictions", "results", "details") if key in value),
            None,
        )
        frame = _door_result_frame(known) if known is not None else pd.DataFrame([value])
    elif isinstance(value, (str, Path)) and Path(value).is_file():
        frame = pd.read_csv(value)
    elif isinstance(value, list):
        if all(isinstance(item, dict) for item in value):
            frame = pd.DataFrame(value)
        elif value and all(isinstance(item, pd.DataFrame) for item in value):
            frame = pd.concat(value, ignore_index=True)
        else:
            raise ValueError("The Door adapter did not return a prediction table.")
    else:
        raise ValueError("The Door adapter did not return a prediction table.")

    renamed = {}
    for column in frame.columns:
        token = _header_token(column)
        if token in {"start", "start_time", "segment_start"}:
            renamed[column] = "start_time"
        elif token in {"end", "end_time", "segment_end"}:
            renamed[column] = "end_time"
        elif token in {"status", "prediction", "predicted", "label"}:
            renamed[column] = "prediction"
    frame = frame.rename(columns=renamed)
    required = {"start_time", "end_time", "prediction"}
    if not required.issubset(frame.columns):
        missing = ", ".join(sorted(required - set(frame.columns)))
        raise ValueError(f"Door adapter result is missing: {missing}.")
    return frame


def _door_visual_module(door_module):
    """Return the original Door signal-processing module for plotting."""
    required = ("load_stream", "split_segments")
    if all(callable(getattr(door_module, name, None)) for name in required):
        return door_module
    try:
        return _load_original_module(
            DOOR_CODE_DIR,
            ("door_standalone.py", "door_standalone(1).py"),
            "door_visual",
            keywords=("door_standalone",),
            extra_dirs=(DOOR_ROOT, APP_DIR),
        )
    except (AttributeError, FileNotFoundError, ImportError, OSError):
        return None


def _door_segment_frames(door_module, path: Path) -> list[tuple[str, pd.DataFrame]]:
    """Build the four traces used by the Door waveform comparison figure."""
    visual_module = _door_visual_module(door_module)
    if visual_module is None:
        return []
    load_stream = getattr(visual_module, "load_stream", None)
    split_segments = getattr(visual_module, "split_segments", None)
    derive = getattr(visual_module, "derive", None)
    if not callable(load_stream) or not callable(split_segments):
        return []

    stream = load_stream(path)
    segments = split_segments(stream)
    output = []
    for segment in segments:
        derived = derive(segment) if callable(derive) else None
        if derived is not None:
            position = pd.to_numeric(derived.get("P"), errors="coerce")
            current = pd.to_numeric(derived.get("I"), errors="coerce")
            voltage = pd.to_numeric(derived.get("V"), errors="coerce")
            back_emf = pd.to_numeric(derived.get("E"), errors="coerce")
            speed = pd.to_numeric(derived.get("speed"), errors="coerce")
            operation = str(derived.attrs.get("op", ""))
        else:
            position = pd.to_numeric(segment["Door leaf position"], errors="coerce")
            current = pd.to_numeric(segment["Motor current(mA)"], errors="coerce")
            voltage = pd.to_numeric(segment["Motor Voltage(10mV)"], errors="coerce")
            back_emf = pd.to_numeric(segment["Motor electrodynamic force"], errors="coerce")
            position_values = position.to_numpy(dtype=float)
            speed_values = (
                np.abs(np.gradient(position_values, 0.02))
                if len(position_values) > 1
                else np.zeros(len(position_values))
            )
            speed = pd.Series(speed_values, index=position.index)
            opening = pd.to_numeric(segment["Door is opening"], errors="coerce")
            operation = "Open" if float(opening.mean()) >= 0.5 else "Close"

        frame = pd.DataFrame(
            {
                "Door leaf position": position,
                "Motor current (mA)": current,
                "Motor voltage (10mV)": voltage,
                "Back-EMF": back_emf,
                "Door speed (pos units/s)": speed,
            }
        ).dropna()
        if not frame.empty:
            output.append((operation or "Unknown", frame))
    return output


def _door_visual_data(door_module, paths: list[Path], names: list[str], details: pd.DataFrame) -> pd.DataFrame:
    """Align Door predictions with original signal segments for the figure."""
    records = []
    offset = 0
    for path, file_id in zip(paths, names):
        segments = _door_segment_frames(door_module, path)
        prediction_rows = details.iloc[offset: offset + len(segments)]
        offset += len(segments)
        for cycle, (operation, frame) in enumerate(segments, start=1):
            prediction = (
                str(prediction_rows.iloc[cycle - 1]["prediction"])
                if cycle <= len(prediction_rows)
                else "Unknown"
            )
            values = frame.copy()
            values["Cycle"] = f"{Path(file_id).stem} · {cycle}"
            values["File"] = file_id
            values["Operation"] = operation
            values["Status"] = prediction
            records.append(values)
    if not records:
        return pd.DataFrame()
    return pd.concat(records, ignore_index=True)


def _door_stream_items(files: list[dict]) -> list[dict]:
    """Return Door inference streams; labels are optional and ignored."""
    streams = []
    for item in files:
        role = (item.get("features") or {}).get("role")
        stem = Path(item["name"]).stem.lower()
        if role == "answer" or "segments_answer" in stem or "answer" in stem:
            continue
        streams.append(item)
    return streams


def files_for_run(subsystem: str, ready: list[dict]) -> list[dict]:
    if subsystem == "Door":
        return _door_stream_items(ready)
    return ready


def can_run(subsystem: str, ready: list[dict]) -> bool:
    if not ready:
        return False
    if subsystem != "Door":
        return True
    return bool(_door_stream_items(ready))


def run_door_prediction(files: list[dict]):
    """Run Door inference on uploaded Train/Test streams only.

    Training labels are intentionally not part of this path.  The production
    adapter/model stored by the Door subsystem is responsible for loading its
    saved parameters.
    """
    if not files:
        raise ValueError("Upload at least one Door Train.csv or Test.csv stream.")
    door = _load_original_module(
        DOOR_CODE_DIR,
        (
            "adapter.py",
            "predict.py",
            "inference.py",
            "door_predict.py",
            "door_standalone.py",
        ),
        "door_predict",
        keywords=("adapter", "predict", "inference", "door"),
        extra_dirs=(DOOR_ROOT,),
    )

    with tempfile.TemporaryDirectory(prefix="door_prediction_") as temp_dir:
        paths = _write_inputs(files, Path(temp_dir))
        if _door_low_level_api(door):
            raise ValueError(
                "The Door module exposes only the development train/answer API. "
                "Use the saved-model adapter.py or production predict.py for deployment inference."
            )

        adapter = _module_entrypoint(
            door,
            (
                "predict_files", "predict_paths", "predict_file", "predict",
                "run_prediction", "analyze_files", "analyze_file", "analyze",
                "infer", "inference", "run",
            ),
        )
        if adapter is None:
            raise AttributeError(
                "The Door subsystem has no saved-model inference entrypoint "
                "(expected adapter.py or predict.py)."
            )
        details = _door_result_frame(_invoke_door_stream_entrypoint(adapter, paths))
        try:
            visual = _door_visual_data(door, paths, [item["name"] for item in files], details)
        except Exception:  # noqa: BLE001
            # Plot preparation must never make a valid Door prediction fail.
            visual = pd.DataFrame()

    submission = details[["start_time", "end_time", "prediction"]].copy()
    return {"details": details, "submission": submission, "visual": visual}


def run_rail_prediction(files: list[dict]):
    rail_predict = _load_original_module(
        RAIL_CODE_DIR,
        ("predict.py",),
        "rail_predict",
        keywords=("predict",),
        extra_dirs=(RAIL_ROOT,),
    )

    with tempfile.TemporaryDirectory(prefix="rail_prediction_") as temp_dir:
        input_paths = _write_inputs(files, Path(temp_dir))
        bundle = rail_predict.load_bundle(RAIL_MODEL_PATH)
        details = rail_predict.predict_files(input_paths, bundle)

    submission = details[["file_id", "prediction"]].copy()
    features = {item["name"]: rail_features(item["key"], item["file"]) for item in files}
    return {"details": details, "submission": submission, "features": features}


def _acv_indoor_timeline(data: bytes) -> pd.DataFrame | None:
    """Indoor average temperature per car over time, for display only."""
    try:
        def wanted(column) -> bool:
            text = str(column)
            match = CAR_COL.match(text)
            if match:
                param = match.group(2).lower()
                return "indoor" in param and "temp" in param
            return "time" in text.lower()

        frame = pd.read_excel(io.BytesIO(data), sheet_name=0, usecols=wanted)
    except Exception:  # noqa: BLE001
        return None
    indoor = {}
    for column in frame.columns:
        match = CAR_COL.match(str(column))
        if match and "indoor" in match.group(2).lower() and "temp" in match.group(2).lower():
            indoor.setdefault(match.group(1), column)
    if not indoor:
        return None

    time_cols = [c for c in frame.columns if "time" in str(c).lower() and not CAR_COL.match(str(c))]
    if time_cols:
        time_axis = pd.to_datetime(frame[time_cols[0]], errors="coerce")
        if time_axis.isna().mean() > 0.5:
            time_axis = pd.Series(np.arange(len(frame)))
    else:
        time_axis = pd.Series(np.arange(len(frame)))

    wide = pd.DataFrame({car: pd.to_numeric(frame[col], errors="coerce") for car, col in indoor.items()})
    wide["time"] = time_axis.values
    step = max(1, len(wide) // 1500)
    wide = wide.iloc[::step]
    return wide.melt(id_vars="time", var_name="car", value_name="indoor_temp_C").dropna()


def run_acv_prediction(files: list[dict]):
    acv_rank = _load_original_module(
        ACV_CODE_DIR,
        ("acv_rank.py", "predict.py", "adapter.py"),
        "acv_rank",
        keywords=("acv", "rank", "predict", "adapter"),
        extra_dirs=(ACV_ROOT,),
    )

    submission_rows, detail_rows, timelines = [], [], {}
    with tempfile.TemporaryDirectory(prefix="acv_prediction_") as temp_dir:
        for item, input_path in zip(files, _write_inputs(files, Path(temp_dir))):
            try:
                result = acv_rank.rank_workbook(input_path)
            except Exception as exc:  # noqa: BLE001
                raise ValueError(f"{item['name']}: {exc}") from exc

            ranked_ids = result.ranked_car_ids  # original header format, e.g. "03"
            submission_rows.append({"file_id": item["name"], "ranked_cars": "|".join(ranked_ids)})
            rank_of = {car_id: pos for pos, car_id in enumerate(ranked_ids, start=1)}
            for i, car_id in enumerate(result.car_ids):
                detail_rows.append(
                    {
                        "file_id": item["name"],
                        "car": car_id,
                        "rank": rank_of[car_id],
                        # The original score is on a 2× temperature scale; halve to get °C.
                        "mean_excess_temp_C": result.scores[i] / 2,
                        "eligible_samples": result.valid_counts[i],
                    }
                )
            timelines[item["name"]] = _acv_indoor_timeline(item["file"].getvalue())

    submission = pd.DataFrame(submission_rows, columns=["file_id", "ranked_cars"])
    details = pd.DataFrame(detail_rows).sort_values(["file_id", "rank"])
    return {"details": details, "submission": submission, "timelines": timelines}


def _read_shm_series(path: Path) -> np.ndarray:
    if path.suffix.lower() == ".npy":
        values = np.asarray(np.load(path, allow_pickle=False))
        if values.size == 0:
            raise ValueError(f"No numeric stress series found in {path.name}.")
        return values.astype(float)
    if path.suffix.lower() == ".npz":
        archive = np.load(path, allow_pickle=False)
        try:
            if not archive.files:
                raise ValueError(f"No numeric stress series found in {path.name}.")
            return np.asarray(archive[archive.files[0]]).astype(float)
        finally:
            archive.close()

    frame = pd.read_csv(path, header=None, sep=None, engine="python")
    numeric = frame.apply(pd.to_numeric, errors="coerce").dropna(axis=1, how="all")
    numeric = numeric.dropna(axis=0, how="all")
    if numeric.empty:
        raise ValueError(f"No numeric stress series found in {path.name}.")
    values = numeric.to_numpy(dtype=float)
    return values[:, 0] if values.shape[1] == 1 else values


def _shm_entrypoint(module):
    names = (
        # Prefer the public inference API.  The bundled Part 4 module also
        # exposes its lower-level ``analyze(data, filename, model)`` helper;
        # calling that helper as a one-argument predictor causes the deployed
        # app to fail with missing ``filename`` and ``model`` arguments.
        "predict_shm", "predict_shm_batch",
        "predict_files", "predict_paths", "predict_file", "predict",
        "run_prediction", "analyze_files", "analyze", "calculate_damage",
        "compute_damage", "predict_damage", "fatigue_damage",
        "run_shm", "infer", "inference", "process_file", "process", "run", "main",
    )
    for name in names:
        function = getattr(module, name, None)
        if callable(function):
            return function
    for name, function in sorted(vars(module).items()):
        lowered = name.lower()
        if (
            callable(function)
            and not name.startswith("_")
            and any(token in lowered for token in ("predict", "infer", "analy", "damage", "fatigue"))
        ):
            return function
    raise AttributeError(
        "The SHM subsystem does not expose a supported prediction entrypoint "
        "(expected predict_files/predict/analyze/calculate_damage or equivalent)."
    )


def _invoke_shm_entrypoint(function, paths: list[Path]):
    """Adapt common original SHM function signatures without changing the module."""
    try:
        parameters = list(inspect.signature(function).parameters.values())
    except (TypeError, ValueError):
        parameters = []
    positional = [
        parameter for parameter in parameters
        if parameter.kind in (parameter.POSITIONAL_ONLY, parameter.POSITIONAL_OR_KEYWORD)
    ]
    required = [parameter for parameter in positional if parameter.default is parameter.empty]
    name = getattr(function, "__name__", "").lower()

    if not required:
        return function()

    # Part 4's batch API expects (filename, source) pairs, not bare paths.
    # Keep this special case in the adapter so the original SHM module remains
    # untouched and its filename validation is preserved. Request its detail
    # fields because the SHM result view uses the waveform and damage bins.
    if name == "predict_shm_batch":
        return function([(path.name, path) for path in paths], include_details=True)

    first = required[0].name.lower()
    wants_files = (
        first in {"files", "paths", "file_paths", "filepaths", "inputs", "datasets"}
        or ("file" in first and first.endswith("s"))
        or ("path" in first and first.endswith("s"))
        or name.endswith("files")
        or name.endswith("paths")
    )
    # Keep the uploaded path opaque unless the original function explicitly
    # asks for a numeric series.  Generic names such as ``data`` often mean a
    # file path in the subsystem adapter; eagerly parsing them here caused the
    # UI to reject otherwise valid SHM files.
    wants_series = any(token in first for token in ("stress", "signal", "series", "array", "values"))

    def invoke_single(path: Path):
        value = _read_shm_series(path) if wants_series else path
        if name == "predict_shm":
            return function(value, include_details=True)
        return function(value)

    if wants_files:
        return function(paths)
    if len(required) > 1:
        # A small adapter commonly exposes predict(train_path, test_path) or
        # predict(file_path, model_path).  Preserve the original function and
        # pass the uploaded paths in order when its arity matches the files.
        if len(required) == len(paths):
            return function(*paths)
        if len(paths) == 1:
            return invoke_single(paths[0])
    if len(paths) == 1:
        return invoke_single(paths[0])

    return [invoke_single(path) for path in paths]


def _normalise_shm_result(value, names: list[str]) -> pd.DataFrame:
    """Turn common SHM return values into one row per uploaded file."""
    if isinstance(value, tuple):
        for part in value:
            if isinstance(part, (pd.DataFrame, dict, list, tuple, np.ndarray)):
                try:
                    return _normalise_shm_result(part, names)
                except ValueError:
                    continue

    if isinstance(value, pd.DataFrame):
        frame = value.copy()
    elif isinstance(value, dict):
        known = next((value[key] for key in ("details", "predictions", "results", "damage") if key in value), None)
        if known is not None and known is not value:
            return _normalise_shm_result(known, names)
        if value and all(str(key) in names or Path(str(key)).name in names for key in value):
            frame = pd.DataFrame({"file_id": [Path(str(key)).name for key in value], "fatigue_damage": list(value.values())})
        else:
            frame = pd.DataFrame([value])
    elif isinstance(value, np.ndarray):
        array = np.asarray(value).reshape(-1)
        frame = pd.DataFrame({"fatigue_damage": array})
    elif isinstance(value, (list, tuple)):
        if value and all(isinstance(item, dict) for item in value):
            frame = pd.DataFrame(value)
        else:
            frame = pd.DataFrame({"fatigue_damage": list(value)})
    else:
        frame = pd.DataFrame({"fatigue_damage": [value]})

    if "file_id" not in frame.columns:
        if len(frame) == len(names):
            frame.insert(0, "file_id", names)
        elif len(names) == 1:
            frame.insert(0, "file_id", names[0])
        else:
            raise ValueError(
                f"SHM returned {len(frame)} result rows for {len(names)} uploaded files."
            )
    else:
        frame["file_id"] = frame["file_id"].map(lambda value: Path(str(value)).name)

    if not any("damage" in str(column).lower() for column in frame.columns):
        prediction_column = next(
            (
                column for column in frame.columns
                if _header_token(column) in {"prediction", "predicted", "damage_prediction"}
            ),
            None,
        )
        if prediction_column is not None:
            frame = frame.rename(columns={prediction_column: "fatigue_damage"})

    if not any("damage" in str(column).lower() for column in frame.columns):
        numeric = [
            column for column in frame.columns
            if column != "file_id" and pd.api.types.is_numeric_dtype(frame[column])
        ]
        if len(numeric) == 1:
            frame = frame.rename(columns={numeric[0]: "fatigue_damage"})
        else:
            raise ValueError("SHM result does not contain a fatigue-damage value column.")
    return frame


def run_shm_prediction(files: list[dict]):
    shm = _load_original_module(
        SHM_CODE_DIR,
        (
            "predict.py",
            "shm_predict.py",
            "shm.py",
            "fatigue.py",
            "damage.py",
            "predict_damage.py",
            "damage_predict.py",
            "inference.py",
            "adapter.py",
            "main.py",
        ),
        "shm_predict",
        keywords=("predict", "shm", "fatigue", "damage", "adapter", "infer"),
        extra_dirs=(SHM_ROOT,),
    )
    function = _shm_entrypoint(shm)
    names = [item["name"] for item in files]
    errors = []
    with tempfile.TemporaryDirectory(prefix="shm_prediction_") as temp_dir:
        paths = _write_inputs(files, Path(temp_dir))
        try:
            # Let an explicitly plural original entrypoint handle the whole
            # upload.  For single-file adapters, _invoke_shm_entrypoint calls
            # one file at a time.
            raw_result = _invoke_shm_entrypoint(function, paths)
            details = _normalise_shm_result(raw_result, names)
        except Exception as batch_error:  # noqa: BLE001
            # One malformed/variant file must not discard the other valid
            # files.  Retry each file independently and keep per-file errors.
            details_parts = []
            for item, path in zip(files, paths):
                try:
                    raw_one = _invoke_shm_entrypoint(function, [path])
                    details_parts.append(_normalise_shm_result(raw_one, [item["name"]]))
                except Exception as exc:  # noqa: BLE001
                    errors.append({"file_id": item["name"], "error": str(exc)})
            if not details_parts:
                raise ValueError(
                    "SHM analysis failed for every uploaded file. "
                    f"First error: {batch_error}"
                ) from batch_error
            details = pd.concat(details_parts, ignore_index=True)

    damage_column = next(column for column in details.columns if "damage" in str(column).lower())
    submission = details[["file_id", damage_column]].copy()
    return {"details": details, "submission": submission, "errors": errors}


RUNNERS = {
    "Door": run_door_prediction,
    "ACV": run_acv_prediction,
    "Rail Corrugation": run_rail_prediction,
    "SHM": run_shm_prediction,
}
OUTPUT_NAMES = {
    "Door": "door_predictions.csv",
    "ACV": "acv_predictions.csv",
    "Rail Corrugation": "rail_predictions.csv",
    "SHM": "shm_predictions.csv",
}


# ---------------------------------------------------------------------------
# Result views
# ---------------------------------------------------------------------------
def verdict(text: str, kind: str) -> None:
    st.markdown(f'<div class="verdict {kind}">{html.escape(text)}</div>', unsafe_allow_html=True)


def download(subsystem: str, submission: pd.DataFrame) -> None:
    st.download_button(
        f"Download {OUTPUT_NAMES[subsystem]}",
        data=submission.to_csv(index=False).encode("utf-8"),
        file_name=OUTPUT_NAMES[subsystem],
        mime="text/csv",
        **FULL,
        help="The file in the official submission format.",
    )


def render_rail(result: dict) -> None:
    details, submission, features = result["details"], result["submission"], result["features"]
    table = submission.copy()
    table["Result"] = table["prediction"].map(RAIL_LABELS).fillna(table["prediction"])
    table["Estimated speed (km/h)"] = table["file_id"].map(
        lambda f: (features.get(f) or {}).get("speed_kmh")
    )
    prob_cols = [c for c in ("p_side1_fault", "p_side2_fault") if c in details.columns]
    if prob_cols:
        table = table.merge(details[["file_id", *prob_cols]], on="file_id", how="left")

    counts = table["prediction"].value_counts()
    n_total = len(table)
    n_fault = n_total - int(counts.get("Normal", 0))

    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Recordings analysed", n_total)
    c2.metric("Both rails normal", int(counts.get("Normal", 0)))
    c3.metric("Side I corrugation", int(counts.get("Side I", 0)))
    c4.metric("Side II corrugation", int(counts.get("Side II", 0)))

    if n_fault == 0:
        verdict("No rail corrugation was detected in any uploaded recording.", "ok")
    else:
        verdict(
            f"{n_fault} of {n_total} recordings show signs of rail corrugation. "
            "The track sections where they were recorded may need inspection.",
            "alert",
        )

    tab_overview, tab_file, tab_table = st.tabs(
        ["Overview", "Inspect a recording", "Full results"]
    )

    with tab_overview:
        left, right = st.columns([1, 1.2], gap="large")
        with left:
            st.markdown("**Results by class**")
            count_df = pd.DataFrame(
                {"Class": list(RAIL_LABELS), "Recordings": [int(counts.get(k, 0)) for k in RAIL_LABELS]}
            )
            chart = (
                alt.Chart(count_df)
                .mark_bar(cornerRadiusEnd=4)
                .encode(
                    y=alt.Y("Class:N", sort=list(RAIL_LABELS), title=None),
                    x=alt.X("Recordings:Q", title="Number of recordings", axis=alt.Axis(tickMinStep=1)),
                    color=alt.Color(
                        "Class:N",
                        scale=alt.Scale(domain=list(RAIL_COLORS), range=list(RAIL_COLORS.values())),
                        legend=None,
                    ),
                    tooltip=["Class", "Recordings"],
                )
                .properties(height=170)
            )
            st.altair_chart(chart, **CHART_FULL)

        with right:
            st.markdown("**Vibration on each rail side**")
            points = []
            for file_id, pred in zip(table["file_id"], table["prediction"]):
                side_rms = (features.get(file_id) or {}).get("side_rms")
                if side_rms:
                    arr = np.asarray(side_rms)
                    points.append(
                        {"File": file_id, "Class": pred,
                         "Side I": float(arr[:, 0].mean()), "Side II": float(arr[:, 1].mean())}
                    )
            if points:
                pts = pd.DataFrame(points)
                upper = float(max(pts["Side I"].max(), pts["Side II"].max())) * 1.05
                diagonal = alt.Chart(pd.DataFrame({"x": [0, upper], "y": [0, upper]})).mark_line(
                    strokeDash=[4, 4], color="#9aa9b6"
                ).encode(x="x:Q", y="y:Q")
                scatter = alt.Chart(pts).mark_circle(size=70, opacity=0.85).encode(
                    x=alt.X("Side I:Q", title="Side I average vibration (m/s²)"),
                    y=alt.Y("Side II:Q", title="Side II average vibration (m/s²)"),
                    color=alt.Color(
                        "Class:N",
                        scale=alt.Scale(domain=list(RAIL_COLORS), range=list(RAIL_COLORS.values())),
                        legend=alt.Legend(orient="bottom", title=None),
                    ),
                    tooltip=["File", "Class",
                             alt.Tooltip("Side I:Q", format=".2f"), alt.Tooltip("Side II:Q", format=".2f")],
                )
                st.altair_chart((diagonal + scatter).properties(height=260), **CHART_FULL)
                st.caption(
                    "Each dot is one recording. Dots above the dashed line vibrate more on "
                    "Side II; dots below it vibrate more on Side I. Vibration also rises with "
                    "train speed, so the model's result takes more into account than this chart."
                )

        flagged = table[table["prediction"] != "Normal"]
        if not flagged.empty:
            st.markdown("**Recordings that need attention**")
            show = ["file_id", "Result", "Estimated speed (km/h)", *prob_cols]
            st.dataframe(
                flagged[show],
                hide_index=True,
                **FULL,
                height=min(38 + 35 * len(flagged), 320),
                column_config={
                    "file_id": st.column_config.TextColumn("File"),
                    "Estimated speed (km/h)": st.column_config.NumberColumn(format="%.0f"),
                    "p_side1_fault": st.column_config.ProgressColumn(
                        "Side I fault likelihood", min_value=0.0, max_value=1.0, format="%.2f"
                    ),
                    "p_side2_fault": st.column_config.ProgressColumn(
                        "Side II fault likelihood", min_value=0.0, max_value=1.0, format="%.2f"
                    ),
                },
            )

    with tab_file:
        order = list(table.sort_values("prediction", key=lambda s: s.eq("Normal"))["file_id"])
        chosen = st.selectbox("Recording", order, help="Recordings with suspected corrugation are listed first.")
        row = table[table["file_id"] == chosen].iloc[0]
        verdict(f"{chosen}: {row['Result']}", "ok" if row["prediction"] == "Normal" else "alert")

        feat = features.get(chosen) or {}
        side_rms = feat.get("side_rms")
        m1, m2, m3 = st.columns(3)
        speed = feat.get("speed_kmh")
        m1.metric("Estimated train speed", f"{speed:.0f} km/h" if speed is not None else "n/a",
                  help="Derived from the speed-sensor pulses (90 teeth, 0.85 m wheel).")
        if side_rms:
            arr = np.asarray(side_rms)
            m2.metric("Side I average vibration", f"{arr[:, 0].mean():.2f} m/s²")
            m3.metric("Side II average vibration", f"{arr[:, 1].mean():.2f} m/s²")

            heat = pd.DataFrame(
                [{"Car": f"Car {c + 1}", "Rail side": side, "Vibration (m/s²)": float(arr[c, s])}
                 for c in range(8) for s, side in enumerate(["Side I", "Side II"])]
            )
            base = alt.Chart(heat).encode(
                x=alt.X("Car:N", sort=[f"Car {i}" for i in range(1, 9)], title=None,
                        axis=alt.Axis(labelAngle=0, orient="top")),
                y=alt.Y("Rail side:N", title=None),
            )
            cells = base.mark_rect(cornerRadius=3).encode(
                color=alt.Color("Vibration (m/s²):Q", scale=alt.Scale(scheme="orangered"),
                                legend=alt.Legend(orient="bottom", title="Vibration (m/s²)")),
                tooltip=["Car", "Rail side", alt.Tooltip("Vibration (m/s²):Q", format=".2f")],
            )
            labels = base.mark_text(fontSize=12).encode(
                text=alt.Text("Vibration (m/s²):Q", format=".1f"),
                color=alt.value("#1b1b1b"),
            )
            st.altair_chart((cells + labels).properties(height=170), **CHART_FULL)
            st.caption(
                "Each cell is the average vibration of the four axle boxes on one side of one car "
                "(8 cars along the train). Corrugation on a rail tends to raise vibration on that "
                "side for many cars, so a row that is darker across most cars is a warning sign."
            )

    with tab_table:
        st.dataframe(details, hide_index=True, **FULL, height=360)

    download("Rail Corrugation", submission)


def render_acv(result: dict) -> None:
    details, submission, timelines = result["details"], result["submission"], result["timelines"]
    files = list(submission["file_id"])

    tab_case, tab_table = st.tabs(["Case result", "Full results"])
    with tab_case:
        chosen = files[0] if len(files) == 1 else st.selectbox("Train case", files)
        case = details[details["file_id"] == chosen].sort_values("rank").reset_index(drop=True)
        top = case.iloc[0]
        verdict(f"Car {top['car']} is the most likely car to have a refrigerant leak.", "alert")

        m1, m2, m3 = st.columns(3)
        m1.metric("Cars compared", len(case))
        m2.metric(f"Car {top['car']} excess temperature", f"{top['mean_excess_temp_C']:.2f} °C")
        if len(case) > 1:
            second = case.iloc[1]
            gap = top["mean_excess_temp_C"] - second["mean_excess_temp_C"]
            m3.metric(f"Lead over Car {second['car']}", f"{gap:.2f} °C",
                      help="The smaller this lead, the more worthwhile it is to inspect the next-ranked cars as well.")

        case["Group"] = np.select(
            [case["rank"] == 1, case["rank"] <= 3], ["Most likely", "Next to check"], "Less likely"
        )
        case["Car label"] = "Car " + case["car"].astype(str)
        bars = (
            alt.Chart(case)
            .mark_bar(cornerRadiusEnd=4)
            .encode(
                y=alt.Y("Car label:N", sort=list(case["Car label"]), title=None),
                x=alt.X("mean_excess_temp_C:Q", title="Average excess temperature (°C)"),
                color=alt.Color(
                    "Group:N",
                    scale=alt.Scale(domain=["Most likely", "Next to check", "Less likely"],
                                    range=["#c0392b", "#e08a00", "#9aa9b6"]),
                    legend=alt.Legend(orient="bottom", title=None),
                ),
                tooltip=[
                    alt.Tooltip("Car label:N", title="Car"),
                    alt.Tooltip("rank:Q", title="Rank"),
                    alt.Tooltip("mean_excess_temp_C:Q", title="Excess temperature (°C)", format=".2f"),
                    alt.Tooltip("eligible_samples:Q", title="Usable data points"),
                ],
            )
            .properties(height=max(180, 30 * len(case)))
        )
        st.altair_chart(bars, **CHART_FULL)
        st.caption(
            "Cars are ordered from most to least likely to be leaking. A higher excess temperature "
            "means the car's cabin stayed warmer than the ranking model expected, which points to "
            "weaker cooling."
        )

        max_samples = case["eligible_samples"].max()
        thin = case[case["eligible_samples"] < 0.1 * max_samples]
        if max_samples > 0 and not thin.empty:
            cars = ", ".join(f"Car {c}" for c in thin["car"])
            st.info(f"{cars}: few usable data points in this case, so the position in the ranking is less reliable.")

        timeline = timelines.get(chosen)
        if timeline is not None and not timeline.empty:
            st.markdown("**Cabin temperature over time**")
            timeline = timeline.assign(
                Highlight=np.where(timeline["car"] == top["car"], f"Car {top['car']}", "Other cars")
            )
            is_time = np.issubdtype(timeline["time"].dtype, np.datetime64)
            lines = alt.Chart(timeline).mark_line(strokeWidth=1.4).encode(
                x=alt.X("time:T" if is_time else "time:Q", title="Time" if is_time else "Sample"),
                y=alt.Y("indoor_temp_C:Q", title="Indoor average temperature (°C)",
                        scale=alt.Scale(zero=False)),
                detail="car:N",
                color=alt.Color("Highlight:N",
                                scale=alt.Scale(domain=[f"Car {top['car']}", "Other cars"],
                                                range=["#c0392b", "#c3cdd6"]),
                                legend=alt.Legend(orient="bottom", title=None)),
                tooltip=[alt.Tooltip("car:N", title="Car"),
                         alt.Tooltip("indoor_temp_C:Q", title="°C", format=".1f")],
            )
            st.altair_chart(lines.properties(height=260), **CHART_FULL)
            st.caption("A leaking unit usually keeps its cabin warmer than the other cars under the same conditions.")

    with tab_table:
        st.dataframe(submission, hide_index=True, **FULL)
        st.dataframe(details, hide_index=True, **FULL, height=320)

    download("ACV", submission)


def render_door(result: dict) -> None:
    details, submission = result["details"], result["submission"]
    abnormal = int((submission["prediction"] == "Abnormal resistance").sum())
    total = len(submission)
    c1, c2, c3 = st.columns(3)
    c1.metric("Door cycles analysed", total)
    c2.metric("Normal", total - abnormal)
    c3.metric("Abnormal resistance", abnormal)
    if abnormal:
        verdict(f"{abnormal} of {total} detected door cycles show abnormal resistance.", "alert")
    else:
        verdict("No abnormal door resistance was detected in the analysed cycles.", "ok")

    visual = result.get("visual", pd.DataFrame())
    if visual.empty:
        st.warning(
            "Door waveform data was not available for plotting. The prediction table is still shown below."
        )
    else:
        operations = [value for value in ("Close", "Open") if value in set(visual["Operation"])]
        options = ["All", *operations]
        selected_operation = st.selectbox(
            "Show door cycles",
            options,
            index=1 if "Close" in operations else 0,
            help="The figure follows the original Door analysis: signal traces versus door leaf position.",
        )
        plot_data = visual if selected_operation == "All" else visual[visual["Operation"] == selected_operation]
        st.markdown(
            f"**Door signal comparison — {selected_operation} cycles**  "
            "*(blue = Normal, red = Abnormal resistance)*"
        )

        color = alt.Color(
            "Status:N",
            scale=alt.Scale(
                domain=["Normal", "Abnormal resistance", "Unknown"],
                range=["#1f77b4", "#d62728", "#7f7f7f"],
            ),
            legend=alt.Legend(title=None, orient="top-right"),
        )
        chart_specs = [
            ("Motor current (mA)", "Motor current (mA)"),
            ("Motor voltage (10mV)", "Motor voltage (10mV)"),
            ("Back-EMF", "Back-EMF"),
            ("Door speed (pos units/s)", "Door speed (pos units/s)"),
        ]
        chart_columns = st.columns(2)
        for index, (field, title) in enumerate(chart_specs):
            with chart_columns[index % 2]:
                chart = (
                    alt.Chart(plot_data)
                    .mark_line(strokeWidth=1.1, opacity=0.78)
                    .encode(
                        x=alt.X(
                            "Door leaf position:Q",
                            title="Door leaf position (0=closed, 700=open)",
                        ),
                        y=alt.Y(f"{field}:Q", title=title),
                        color=color,
                        detail="Cycle:N",
                        tooltip=[
                            alt.Tooltip("Cycle:N"),
                            alt.Tooltip("Operation:N"),
                            alt.Tooltip("Status:N"),
                            alt.Tooltip("Door leaf position:Q", format=".1f"),
                            alt.Tooltip(f"{field}:Q", title=title, format=".4g"),
                        ],
                    )
                    .properties(height=280)
                )
                st.altair_chart(chart, **CHART_FULL)

    st.dataframe(submission, hide_index=True, **FULL, height=360)
    if "p_abnormal" in details.columns:
        st.markdown("**Model details**")
        st.dataframe(details, hide_index=True, **FULL, height=360)
    download("Door", submission)


def _shm_waveform_frame(value) -> pd.DataFrame:
    """Convert optional SHM chart points into a safe plotting table."""
    if isinstance(value, np.ndarray):
        value = value.tolist()
    if not isinstance(value, (list, tuple)):
        return pd.DataFrame(columns=["Sample", "Stress"])

    points = []
    for item in value:
        if isinstance(item, dict):
            sample = item.get("sample_index", item.get("sample", item.get("x")))
            stress = item.get("stress", item.get("value", item.get("y")))
        elif isinstance(item, (list, tuple, np.ndarray)) and len(item) >= 2:
            sample, stress = item[0], item[1]
        else:
            continue
        try:
            points.append({"Sample": float(sample), "Stress": float(stress)})
        except (TypeError, ValueError):
            continue
    return pd.DataFrame(points, columns=["Sample", "Stress"])


def _shm_bins_frame(value) -> pd.DataFrame:
    """Convert optional SHM amplitude-bin details into a plotting table."""
    if isinstance(value, np.ndarray):
        value = value.tolist()
    if not isinstance(value, (list, tuple)):
        return pd.DataFrame(columns=["low", "high", "count", "damage"])

    bins = []
    for item in value:
        if not isinstance(item, dict):
            continue
        try:
            bins.append(
                {
                    "low": float(item["low"]),
                    "high": float(item["high"]),
                    "count": float(item.get("count", 0.0)),
                    "damage": float(item.get("damage", 0.0)),
                }
            )
        except (KeyError, TypeError, ValueError):
            continue
    return pd.DataFrame(bins, columns=["low", "high", "count", "damage"])


def render_shm(result: dict) -> None:
    details, submission = result["details"], result["submission"]
    damage_column = next(column for column in submission.columns if column != "file_id")
    values = pd.to_numeric(submission[damage_column], errors="coerce")
    errors = result.get("errors", [])
    if errors:
        st.warning(
            f"{len(errors)} file(s) could not be analysed; successful files are shown below."
        )
        st.dataframe(pd.DataFrame(errors), hide_index=True, **FULL)
    c1, c2, c3 = st.columns(3)
    c1.metric("Segments analysed", len(submission))
    c2.metric("Maximum fatigue damage", f"{values.max():.6g}" if values.notna().any() else "n/a")
    c3.metric("Total fatigue damage", f"{values.sum():.6g}" if values.notna().any() else "n/a")

    overview = submission[["file_id", damage_column]].copy()
    overview["File"] = overview["file_id"].astype(str)
    overview["Damage"] = pd.to_numeric(overview[damage_column], errors="coerce")
    damage_chart = (
        alt.Chart(overview)
        .mark_bar(cornerRadiusEnd=4, color="#2e8b57")
        .encode(
            x=alt.X("File:N", sort="-y", title=None),
            y=alt.Y("Damage:Q", title="Cumulative fatigue damage"),
            tooltip=[
                alt.Tooltip("File:N", title="File"),
                alt.Tooltip("Damage:Q", title="Damage", format=".6g"),
            ],
        )
        .properties(height=260)
    )
    st.markdown("**Damage by uploaded file**")
    st.altair_chart(damage_chart, **CHART_FULL)

    file_names = [str(name) for name in submission["file_id"]]
    selected = st.selectbox("Inspect SHM recording", file_names)
    selected_rows = details[details["file_id"].astype(str) == selected]
    if not selected_rows.empty:
        row = selected_rows.iloc[0]
        warnings = row.get("warnings")
        if isinstance(warnings, str) and warnings:
            warnings = [warnings]
        if isinstance(warnings, (list, tuple, np.ndarray)):
            for warning in warnings:
                if str(warning).strip():
                    st.info(str(warning))

        waveform = _shm_waveform_frame(row.get("waveform"))
        bins = _shm_bins_frame(row.get("bins"))
        if not waveform.empty or not bins.empty:
            tab_wave, tab_bins = st.tabs(["Stress waveform", "Damage by amplitude"])
            with tab_wave:
                if waveform.empty:
                    st.info("This SHM adapter did not return waveform details.")
                else:
                    line = (
                        alt.Chart(waveform)
                        .mark_line(color="#1d6fa5", strokeWidth=1.5)
                        .encode(
                            x=alt.X("Sample:Q", title="Sample index"),
                            y=alt.Y("Stress:Q", title="Stress"),
                            tooltip=[
                                alt.Tooltip("Sample:Q", title="Sample", format=".0f"),
                                alt.Tooltip("Stress:Q", title="Stress", format=".4g"),
                            ],
                        )
                        .properties(height=280)
                    )
                    st.altair_chart(line, **CHART_FULL)
                    st.caption(
                        "The displayed line is an extrema-preserving envelope; the model uses the full uploaded signal."
                    )
            with tab_bins:
                if bins.empty:
                    st.info("This SHM adapter did not return amplitude-bin details.")
                else:
                    bars = (
                        alt.Chart(bins)
                        .mark_bar(color="#e08a00")
                        .encode(
                            x=alt.X("low:Q", title="Cycle amplitude"),
                            x2="high:Q",
                            y=alt.Y("damage:Q", title="Damage contribution"),
                            tooltip=[
                                alt.Tooltip("low:Q", title="Low", format=".4g"),
                                alt.Tooltip("high:Q", title="High", format=".4g"),
                                alt.Tooltip("count:Q", title="Half-cycle count", format=".4g"),
                                alt.Tooltip("damage:Q", title="Damage", format=".6g"),
                            ],
                        )
                        .properties(height=280)
                    )
                    st.altair_chart(bars, **CHART_FULL)

    st.dataframe(submission, hide_index=True, **FULL, height=320)
    if len(details.columns) > len(submission.columns):
        st.markdown("**Model details**")
        detail_table = details.drop(columns=[column for column in ("waveform", "bins") if column in details.columns])
        st.dataframe(detail_table, hide_index=True, **FULL, height=360)
    download("SHM", submission)


RENDERERS = {
    "Door": render_door,
    "ACV": render_acv,
    "Rail Corrugation": render_rail,
    "SHM": render_shm,
}


# ---------------------------------------------------------------------------
# Layout
# ---------------------------------------------------------------------------
st.markdown(
    """
    <div class="hero">
        <h1>Train Condition Monitoring System</h1>
        <p>Choose a subsystem, upload its sensor files and run the analysis.</p>
    </div>
    """,
    unsafe_allow_html=True,
)

subsystem = st.radio("Subsystem", list(SUBSYSTEMS), horizontal=True, label_visibility="collapsed")
spec = SUBSYSTEMS[subsystem]
st.caption(spec["purpose"])

left, right = st.columns([1.25, 0.75], gap="large")

with right:
    st.markdown(
        f"""
        <div class="guide">
            <p><b>What to upload:</b> {html.escape(spec['upload'])}</p>
            <p><b>Expected format:</b> {html.escape(spec['structure'])}</p>
            <p><b>What you get:</b> {html.escape(spec['result'])}</p>
        </div>
        """,
        unsafe_allow_html=True,
    )

with left:
    version = st.session_state["uploader_version"][subsystem]
    uploaded_files = st.file_uploader(
        "Upload files",
        accept_multiple_files=True,
        label_visibility="collapsed",
        key=f"uploader_{subsystem}_{version}",
    )

    checked = check_files(subsystem, uploaded_files) if uploaded_files else []
    ready = [item for item in checked if item["status"] != "error"]
    rejected = [item for item in checked if item["status"] == "error"]
    run_items = []
    run_ready = can_run(subsystem, ready)
    if run_ready:
        run_items = files_for_run(subsystem, ready)

    if checked:
        col_kwargs = (
            {"vertical_alignment": "center"}
            if "vertical_alignment" in inspect.signature(st.columns).parameters
            else {}
        )
        info_col, clear_col = st.columns([3, 1], **col_kwargs)
        summary = f"**{len(checked)}** file(s) selected, **{len(ready)}** ready"
        if rejected:
            summary += f", **{len(rejected)}** not accepted"
        info_col.markdown(summary)
        if clear_col.button("Remove all", **FULL, help="Remove every uploaded file."):
            st.session_state["uploader_version"][subsystem] += 1
            st.rerun()

        status_text = {"ok": "✅ Ready", "warn": "⚠️ Ready, see note", "error": "❌ Not accepted"}
        file_table = pd.DataFrame(
            {
                "File": [item["name"] for item in checked],
                "Size": [human_size(item["size"]) for item in checked],
                "Check": [status_text[item["status"]] for item in checked],
                "Note": [item["note"] for item in checked],
            }
        )
        st.dataframe(
            file_table,
            hide_index=True,
            **FULL,
            height=min(38 + 35 * len(checked), 250),
            column_config={
                "File": st.column_config.TextColumn(width="medium"),
                "Size": st.column_config.TextColumn(width="small"),
                "Check": st.column_config.TextColumn(width="small"),
                "Note": st.column_config.TextColumn(width="large"),
            },
        )
        if rejected:
            st.error(
                f"{len(rejected)} file(s) do not match the {subsystem} format and will be skipped. "
                "The Note column explains each problem. Expected format: " + spec["structure"]
            )

    run_disabled = not spec["connected"] or not run_ready
    run_clicked = st.button("Run analysis", type="primary", **FULL, disabled=run_disabled)
    if not checked:
        st.caption("Upload at least one file to start.")
    elif subsystem == "Door" and not run_ready:
        st.caption("Upload at least one Door Train.csv or Test.csv stream file.")
    elif not ready:
        st.caption("None of the uploaded files can be analysed. Please check the notes above.")


# ---------------------------------------------------------------------------
# Run + results
# ---------------------------------------------------------------------------
if run_clicked:
    try:
        with st.spinner(f"Analysing {len(run_items)} file(s)..."):
            outcome = RUNNERS[subsystem](run_items)
        outcome["inputs"] = sorted(item["name"] for item in run_items)
        st.session_state["results"][subsystem] = outcome
    except Exception as exc:  # noqa: BLE001
        st.error(f"The {subsystem} analysis could not be completed: {exc}")

_fragment = getattr(st, "fragment", None) or (lambda func: func)


@_fragment
def show_results(subsystem: str, current_inputs: list[str]) -> None:
    """Rendered as a fragment: switching tabs or recordings reruns only this block."""
    outcome = st.session_state["results"].get(subsystem)
    if outcome is None:
        return
    st.markdown('<div class="section-title">Results</div>', unsafe_allow_html=True)
    if current_inputs and current_inputs != outcome["inputs"]:
        st.caption("The file list has changed since these results were produced. Run the analysis again to update them.")
    RENDERERS[subsystem](outcome)


show_results(subsystem, sorted(item["name"] for item in run_items))
